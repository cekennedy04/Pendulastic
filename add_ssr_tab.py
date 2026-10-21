"""Add an 'SSR Step-Response' tab to an existing Bench_Experiment_Tracker.xlsx.

Preserves everything already in the workbook -- opens it, appends the one sheet
(idempotent: skips if the sheet already exists), saves. Run it on YOUR current
tracker so your hand-entered rows are kept:

    python add_ssr_tab.py "path/to/Bench_Experiment_Tracker.xlsx"

One row per STEP. The Stretch-Mode staircase (0->+A->-A->+A->0, 3 s dwells,
6 s final) yields four steps per recording; log each. Fill the measured
columns; error, %error and the gain/linearity/drift analysis compute themselves.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import ScatterChart, Reference, Series

PATH = sys.argv[1] if len(sys.argv) > 1 else "Pendulastic_Mobile_App_Trial/Bench_Experiment_Tracker.xlsx"
NROWS = 120

INK="16202B"; HEAD="16202B"; ACC="0D6A86"; ACC_SOFT="DCECF2"; GREY="5C6B78"
COND="EAF0F4"; MEAS="DFEEE6"; COMP="FBEEDD"; RULE="CFDAE2"; RULE_SOFT="E2EAEF"; ZEBRA="F6F9FB"
def F(**k): return Font(name="Calibri", **k)
def fill(c): return PatternFill("solid", fgColor=c)
thin=Side(style="thin", color=RULE); box=Border(left=thin,right=thin,top=thin,bottom=thin)

wb = openpyxl.load_workbook(PATH)
if "SSR Step-Response" in wb.sheetnames:
    print("SSR Step-Response tab already present — nothing to do.")
    sys.exit(0)
ws = wb.create_sheet("SSR Step-Response")

# title
ws["A1"] = "SSR — steady-state step-response angle accuracy"
ws["A1"].font = F(bold=True, size=14, color=INK)
ws["A2"] = ("One row per step. Fill Encoder Δ and IMU Δ (settled-to minus settled-from) and the drift; "
            "error, %err and the gain/linearity below compute themselves. Horizontal rig: read step Δ, not absolute angle.")
ws["A2"].font = F(size=9, italic=True, color=GREY); ws.merge_cells("A2:P2")
ws["A2"].alignment = Alignment(wrap_text=True); ws.row_dimensions[2].height = 30

GROUPS = [
 ("IDENTITY", ACC_SOFT, [
   ("Rec #","the recording this step is in",7),
   ("Capture rank","true run order (bench_rescore --capture-order)",10),
   ("Encoder file","the .csv it maps to",26),("Date","",11)]),
 ("CONDITION", COND, [
   ("P","P term",6),("A (deg)","Stretch Angle Range",8),("Step speed","°/s",9),
   ("Step #","1-4 within the recording",7),
   ("From°","commanded start angle",7),("To°","commanded end angle",7),
   ("Cmd Δ","=To−From",8)]),
 ("MEASURED", MEAS, [
   ("Encoder Δ","settled to − from (deg)",10),
   ("IMU Δ","settled to − from (deg)",10),
   ("Drift °/s","slope of IMU during the 3 s dwell",11)]),
 ("COMPUTED", COMP, [
   ("Δ error","=IMU Δ − Encoder Δ",9),("% error","of encoder Δ",8),("|Δ err|","for the stats",8)]),
 ("", "F0F0F0", [("Notes","",34)]),
]
c=1
for gname,gcol,cols in GROUPS:
    start=c
    for hdr,note,w in cols:
        h=ws.cell(row=4,column=c,value=hdr); h.font=F(bold=True,size=10,color="FFFFFF")
        h.fill=fill(HEAD); h.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); h.border=box
        n=ws.cell(row=5,column=c,value=note); n.font=F(size=8,italic=True,color=GREY)
        n.alignment=Alignment(horizontal="center",vertical="top",wrap_text=True); n.border=box
        ws.column_dimensions[get_column_letter(c)].width=w; c+=1
    if gname:
        ws.merge_cells(start_row=3,start_column=start,end_row=3,end_column=c-1)
        g=ws.cell(row=3,column=start,value=gname); g.font=F(bold=True,size=10,color=INK)
        g.fill=fill(gcol); g.alignment=Alignment(horizontal="center",vertical="center")
ws.row_dimensions[3].height=18; ws.row_dimensions[4].height=26; ws.row_dimensions[5].height=30
ws.freeze_panes="A6"

hdrs=[ws.cell(row=4,column=i).value for i in range(1,c)]
def col(h): return hdrs.index(h)+1
def L(h): return get_column_letter(col(h))

for r in range(6,6+NROWS):
    ws.cell(row=r,column=col("Cmd Δ"),
            value=f'=IF(AND(ISNUMBER({L("From°")}{r}),ISNUMBER({L("To°")}{r})),{L("To°")}{r}-{L("From°")}{r},"")')
    ws.cell(row=r,column=col("Δ error"),
            value=f'=IF(AND(ISNUMBER({L("Encoder Δ")}{r}),ISNUMBER({L("IMU Δ")}{r})),{L("IMU Δ")}{r}-{L("Encoder Δ")}{r},"")')
    pe=ws.cell(row=r,column=col("% error"),
            value=f'=IF(AND(ISNUMBER({L("Encoder Δ")}{r}),{L("Encoder Δ")}{r}<>0),({L("IMU Δ")}{r}-{L("Encoder Δ")}{r})/{L("Encoder Δ")}{r},"")')
    pe.number_format="0.0%"
    ws.cell(row=r,column=col("|Δ err|"),
            value=f'=IF(ISNUMBER({L("Δ error")}{r}),ABS({L("Δ error")}{r}),"")')
    for cc in range(1,c):
        cell=ws.cell(row=r,column=cc); cell.border=box; cell.font=F(size=10)
        if r%2==0: cell.fill=fill(ZEBRA)
        cell.alignment=Alignment(horizontal=("left" if ws.cell(4,cc).value in ("Notes","Encoder file") else "center"),vertical="center")

# dropdown for Step #
dv=DataValidation(type="list",formula1='"1,2,3,4"',allow_blank=True); ws.add_data_validation(dv)
dv.add(f"{L('Step #')}6:{L('Step #')}{5+NROWS}")

# ── ANALYSIS block (formulas over the measured columns) ──
er=6+NROWS+2
ws.cell(row=er,column=1,value="ANALYSIS  (fills as rows are entered)").font=F(bold=True,size=11,color=ACC)
enc=f'{L("Encoder Δ")}6:{L("Encoder Δ")}{5+NROWS}'; imu=f'{L("IMU Δ")}6:{L("IMU Δ")}{5+NROWS}'
absE=f'{L("|Δ err|")}6:{L("|Δ err|")}{5+NROWS}'; dr=f'{L("Drift °/s")}6:{L("Drift °/s")}{5+NROWS}'
# All ordinary (non-array) formulas: AVERAGE/MAX ignore the "" cells, SLOPE/RSQ
# ignore non-numeric pairs. No Ctrl+Shift+Enter needed anywhere.
rows=[
 ("Gain (slope IMU Δ vs Encoder Δ) — ideal 1.000", f'=IFERROR(SLOPE({imu},{enc}),"")'),
 ("Intercept (deg) — ideal 0", f'=IFERROR(INTERCEPT({imu},{enc}),"")'),
 ("Linearity R²  — ideal 1.000", f'=IFERROR(RSQ({imu},{enc}),"")'),
 ("Mean |Δ error| (deg)", f'=IFERROR(AVERAGE({absE}),"")'),
 ("Max |Δ error| (deg)", f'=IFERROR(MAX({absE}),"")'),
 ("Mean drift (°/s)", f'=IFERROR(AVERAGE({dr}),"")'),
 ("Steps logged", f'=COUNT({enc})'),
]
for i,(lab,formula) in enumerate(rows):
    a=ws.cell(row=er+1+i,column=1,value=lab); a.font=F(size=10,bold=("Gain" in lab or "Linearity" in lab)); a.border=box; a.fill=fill(ACC_SOFT)
    b=ws.cell(row=er+1+i,column=2,value=formula); b.border=box; b.font=F(size=10)
    if "Mean |" in lab or "Max |" in lab or "Gain" in lab or "R²" in lab or "drift" in lab or "Intercept" in lab:
        b.number_format="0.000"
ws.merge_cells(start_row=er, start_column=1, end_row=er, end_column=2)

# ── scatter chart: IMU Δ (y) vs Encoder Δ (x) ──
chart=ScatterChart(); chart.title="IMU Δ vs Encoder Δ (points on the 45° line = accurate)"
chart.x_axis.title="Encoder Δ (deg)"; chart.y_axis.title="IMU Δ (deg)"; chart.height=9; chart.width=14
xref=Reference(ws, min_col=col("Encoder Δ"), min_row=6, max_row=5+NROWS)
yref=Reference(ws, min_col=col("IMU Δ"), min_row=6, max_row=5+NROWS)
s=Series(yref, xref, title="steps"); s.marker.symbol="circle"; s.graphicalProperties.line.noFill=True
chart.series.append(s)
ws.add_chart(chart, f"D{er}")

wb.save(PATH)
print(f"Added 'SSR Step-Response' to {PATH}  (existing sheets untouched: {', '.join(n for n in wb.sheetnames if n!='SSR Step-Response')})")
