"""Build the bench experiment tracker workbook.

Sheet-centric workflow (the Bench Runner artifact is retired):
  How to use · Clean Session Plan · Trials · Accuracy · Summary · Protocol · Session

The Trials sheet is the running log; bench_rescore.py writes the ENCODER TRUTH
columns back into it (match on Phone JSONL file). Accuracy and Summary are pure
formulas over Trials, so they update the moment truth lands.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
INK="16202B"; ACC="0D6A86"; ACC_SOFT="DCECF2"; HEAD="16202B"
WARN_SOFT="F6E6DA"; GOOD_SOFT="DFEEE6"; GREY="5C6B78"; RULE="CFDAE2"; ZEBRA="F6F9FB"
def F(**k): return Font(name="Calibri", **k)
def fill(c): return PatternFill("solid", fgColor=c)
thin=Side(style="thin", color=RULE); box=Border(left=thin,right=thin,top=thin,bottom=thin)
NROWS=140

wb=openpyxl.Workbook()

# ════════════════════════════════════════════════════ TRIALS (running log)
ws=wb.active; ws.title="Trials"
GROUPS=[
 ("IDENTITY", ACC_SOFT, [
   ("Trial #","phone app trial number",7),("Session","which session",13),
   ("Phone JSONL file","exact export filename — the key bench_rescore matches on",30),
   ("Encoder file","the .csv stem it maps to",28),
   ("Encoder seg","which ring: 2/4/6/8",10),("Mode","Stretch / Manual",9),("Date","",11),
 ]),
 ("RIG SETTINGS", "EAF0F4", [
   ("P","",6),("I","",6),("D","",8),("D switch","",8),
   ("Filter","No Filter / PID FIR / Butterworth",12),("Output ±","",8),
   ("Trial Type","Fast / Slow",9),("Fast Stretch","lift °/s",10),("Final Return","release °/s",10),
   ("Stretch Range","A0 deg",11),("Prep s","",7),("Hold-out s","",9),("Hold-neu s","",9),
   ("Total Trials","",10),("Lim CCW","",8),("Lim CW","",8),("Zero off°","",9),
 ]),
 ("OBSERVED", GOOD_SOFT, [
   ("Crossings","",9),("Settle s","",8),("f calc","=cross/(2·s)",8),("tau calc","=s/ln20",8),
 ]),
 ("APP READOUT", "FBEEDD", [
   ("app A0","",8),("app A1","",8),("app N","",6),("app f","",7),("app R2n","",7),
   ("app area","",8),("app wpeak","",9),("Refused?","yes/no",9),
 ]),
 ("ENCODER TRUTH  (bench_rescore.py writes here)", "E7E0F0", [
   ("enc A0","",8),("enc N","",6),("enc f","",7),("enc R2n","",7),("enc wmax","",8),("enc area","",8),
 ]),
 ("STATUS", "F0F0F0", [
   ("Status","logged/recorded/evidence/faulted/skipped",11),
   ("Notes / deviations","",42),
 ]),
]
c=1
for gname,gcol,cols in GROUPS:
    start=c
    for hdr,note,w in cols:
        cell=ws.cell(row=2,column=c,value=hdr); cell.font=F(bold=True,size=10,color="FFFFFF")
        cell.fill=fill(HEAD); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=box
        h=ws.cell(row=3,column=c,value=note); h.font=F(size=8,italic=True,color=GREY)
        h.alignment=Alignment(horizontal="center",vertical="top",wrap_text=True); h.border=box
        ws.column_dimensions[get_column_letter(c)].width=w; c+=1
    ws.merge_cells(start_row=1,start_column=start,end_row=1,end_column=c-1)
    g=ws.cell(row=1,column=start,value=gname); g.font=F(bold=True,size=10,color=INK)
    g.fill=fill(gcol); g.alignment=Alignment(horizontal="center",vertical="center")
ws.row_dimensions[1].height=20; ws.row_dimensions[2].height=30; ws.row_dimensions[3].height=32
ws.freeze_panes="H4"
hdrs=[ws.cell(row=2,column=i).value for i in range(1,c)]
def col(h): return hdrs.index(h)+1
def L(h): return get_column_letter(col(h))
for r in range(4,4+NROWS):
    ws.cell(row=r,column=col("f calc"),value=f'=IF(AND(ISNUMBER({L("Crossings")}{r}),{L("Settle s")}{r}>0),{L("Crossings")}{r}/(2*{L("Settle s")}{r}),"")')
    ws.cell(row=r,column=col("tau calc"),value=f'=IF({L("Settle s")}{r}>0,{L("Settle s")}{r}/LN(20),"")')
    for cc in range(1,c):
        cell=ws.cell(row=r,column=cc); cell.border=box; cell.font=F(size=10)
        if r%2==0: cell.fill=fill(ZEBRA)
        cell.alignment=Alignment(horizontal=("left" if ws.cell(row=2,column=cc).value in ("Notes / deviations","Phone JSONL file","Encoder file") else "center"),vertical="center")
def dv(colname, options, sheet=ws, nrows=NROWS, r0=4):
    d=DataValidation(type="list",formula1='"'+",".join(options)+'"',allow_blank=True); sheet.add_data_validation(d)
    d.add(f"{get_column_letter([sheet.cell(row=2,column=i).value for i in range(1,sheet.max_column+1)].index(colname)+1)}{r0}:{get_column_letter([sheet.cell(row=2,column=i).value for i in range(1,sheet.max_column+1)].index(colname)+1)}{r0-1+nrows}")
dv("D switch",["on","off"]); dv("Filter",["No Filter","PID FIR Filter","Butterworth"])
dv("Trial Type",["Fast","Slow"]); dv("Mode",["Stretch","Manual"]); dv("Encoder seg",["2","4","6","8"])
dv("Refused?",["yes","no"]); dv("Status",["logged","recorded","evidence","faulted","skipped"])
TRIALS_MAXROW=3+NROWS

# ════════════════════════════════════════════════════ CLEAN SESSION PLAN
cp=wb.create_sheet("Clean Session Plan"); wb.move_sheet("Clean Session Plan", -(len(wb.sheetnames)-1))
cp["A1"]="Clean Session Plan"; cp["A1"].font=F(bold=True,size=15,color=INK)
cp["A2"]=("Stretch mode only. One phone recording per encoder ring — start recording during the Prepare hold, "
          "stop after the neutral hold. Write the encoder file + segment into the Trials sheet the instant you record. "
          "Everything else stays at the working config unless the row says to change it.")
cp["A2"].font=F(size=9,italic=True,color=GREY); cp.merge_cells("A2:H2"); cp["A2"].alignment=Alignment(wrap_text=True)
cp.row_dimensions[2].height=42
ph=["#","Block","Change THIS","to","Reps","Records what","Encoder file (fill in)","Seg"]
for j,h in enumerate(ph,1):
    cell=cp.cell(row=4,column=j,value=h); cell.font=F(bold=True,color="FFFFFF"); cell.fill=fill(HEAD)
    cell.alignment=Alignment(horizontal="center",wrap_text=True); cell.border=box
PLAN=[]
n=1
for rep in range(1,7):
    PLAN.append((n,"A repeatability","nothing","working config",f"{rep}/6","seg-2 ring, identical","","2")); n+=1
for a0 in (37.5,30,20,15,10,5):
    PLAN.append((n,"B amplitude","Stretch Range",a0,"1","A0 at this amplitude","","2")); n+=1
for pval,hz in ((0.43,"0.6 Hz"),(1.7,"1.2 Hz")):
    PLAN.append((n,"C frequency","P term",f"{pval} ({hz})","1","f at this P","","2")); n+=1
for off in (0,3,6):
    PLAN.append((n,"E asymmetry","ZeroAngle offset",f"{off}°","1","P_asym; 0° = null test","","2")); n+=1
for i,rowv in enumerate(PLAN):
    for j,v in enumerate(rowv,1):
        cell=cp.cell(row=5+i,column=j,value=v); cell.border=box; cell.font=F(size=10)
        cell.alignment=Alignment(horizontal=("left" if j in (6,7) else "center"),vertical="center")
        if i%2: cell.fill=fill(ZEBRA)
tot=cp.cell(row=6+len(PLAN),column=1,value=f"{len(PLAN)} recordings"); tot.font=F(bold=True)
cp.cell(row=6+len(PLAN),column=6,value="working config = P1 · I0 · D off · No Filter · ±50 · Fast · lift 50 · return 200 · prep 3 · holds 15/15").font=F(size=9,italic=True,color=ACC)
for j,w in enumerate([4,16,16,16,6,22,24,6],1): cp.column_dimensions[get_column_letter(j)].width=w
cp.column_dimensions["A"].width=4

# ════════════════════════════════════════════════════ ACCURACY
cs=wb.create_sheet("Accuracy")
cs["A1"]="Phone vs encoder — per-metric error"; cs["A1"].font=F(bold=True,size=14,color=INK)
cs["A2"]="Formulas over the Trials rows. Populates as bench_rescore.py fills the encoder-truth columns."
cs["A2"].font=F(size=9,italic=True,color=GREY)
for j,h in enumerate(["Trial #","Metric","App value","Encoder truth","Error","% of truth","|%|"],1):
    cell=cs.cell(row=4,column=j,value=h); cell.font=F(bold=True,color="FFFFFF"); cell.fill=fill(HEAD)
    cell.alignment=Alignment(horizontal="center"); cell.border=box
metrics=[("A0","app A0","enc A0"),("N","app N","enc N"),("f","app f","enc f"),
         ("R2n","app R2n","enc R2n"),("area","app area","enc area")]
row=5
for i in range(NROWS):
    tr=4+i
    for label,appc,encc in metrics:
        cs.cell(row=row,column=1,value=f'=IF(Trials!{L("Trial #")}{tr}="","",Trials!{L("Trial #")}{tr})')
        cs.cell(row=row,column=2,value=label)
        cs.cell(row=row,column=3,value=f"=Trials!{L(appc)}{tr}")
        cs.cell(row=row,column=4,value=f"=Trials!{L(encc)}{tr}")
        cs.cell(row=row,column=5,value=f'=IF(AND(ISNUMBER(C{row}),ISNUMBER(D{row})),C{row}-D{row},"")')
        cs.cell(row=row,column=6,value=f'=IF(AND(ISNUMBER(C{row}),ISNUMBER(D{row}),D{row}<>0),(C{row}-D{row})/D{row},"")'); cs.cell(row=row,column=6).number_format="0%"
        cs.cell(row=row,column=7,value=f'=IF(ISNUMBER(F{row}),ABS(F{row}),"")'); cs.cell(row=row,column=7).number_format="0%"
        for j in range(1,8): cs.cell(row=row,column=j).border=box; cs.cell(row=row,column=j).alignment=Alignment(horizontal="center")
        row+=1
ACC_MAXROW=row-1
for j,w in enumerate([9,9,12,14,12,11,9],1): cs.column_dimensions[get_column_letter(j)].width=w
cs.freeze_panes="A5"

# ════════════════════════════════════════════════════ SUMMARY (preliminary analysis)
sm=wb.create_sheet("Summary")
sm["A1"]="Preliminary analysis"; sm["A1"].font=F(bold=True,size=15,color=INK)
sm["A2"]="Auto-computed from Trials + Accuracy. Numbers appear as rows fill in."; sm["A2"].font=F(size=9,italic=True,color=GREY)
def hcell(r,cvals,fillc=HEAD,fc="FFFFFF"):
    for j,v in enumerate(cvals,1):
        cell=sm.cell(row=r,column=j,value=v); cell.font=F(bold=True,color=fc); cell.fill=fill(fillc)
        cell.alignment=Alignment(horizontal="center",wrap_text=True); cell.border=box
def datarow(r,cvals):
    for j,v in enumerate(cvals,1):
        cell=sm.cell(row=r,column=j,value=v); cell.border=box; cell.font=F(size=10)
        cell.alignment=Alignment(horizontal="left" if j==1 else "center")

# Coverage
sm.cell(row=4,column=1,value="COVERAGE").font=F(bold=True,size=11,color=ACC)
hcell(5,["Metric","Value"])
St=lambda h: f"Trials!{L(h)}4:{L(h)}{TRIALS_MAXROW}"
cov=[("Phone trials logged",f'=COUNTA({St("Phone JSONL file")})'),
     ("Recorded (has encoder file)",f'=COUNTIF({St("Encoder file")},"?*")'),
     ("App refused (unmeasured)",f'=COUNTIF({St("Refused?")},"yes")'),
     ("Refusal rate",f'=IFERROR(COUNTIF({St("Refused?")},"yes")/COUNTA({St("Refused?")}),"")'),
     ("Encoder truth filled (enc f)",f'=COUNT({St("enc f")})')]
for i,(k,v) in enumerate(cov):
    datarow(6+i,[k,v])
    if "rate" in k: sm.cell(row=6+i,column=2).number_format="0%"

# Per-metric error, from Accuracy sheet
base=13
sm.cell(row=base,column=1,value="PER-METRIC ERROR  (app − encoder, over all paired trials)").font=F(bold=True,size=11,color=ACC)
hcell(base+1,["Metric","n pairs","mean app","mean truth","mean |%err|","pass? (see note)"])
AccM=f"Accuracy!$B$5:$B${ACC_MAXROW}"; AccApp=f"Accuracy!$C$5:$C${ACC_MAXROW}"; AccEnc=f"Accuracy!$D$5:$D${ACC_MAXROW}"; AccAbs=f"Accuracy!$G$5:$G${ACC_MAXROW}"
thresh={"f":0.05,"A0":0.10,"N":0.15,"R2n":0.15,"area":0.20}
for i,(m,_,_) in enumerate(metrics):
    r=base+2+i
    datarow(r,[m,
        f'=COUNTIFS({AccM},"{m}",{AccApp},"<>",{AccEnc},"<>")',
        f'=IFERROR(AVERAGEIFS({AccApp},{AccM},"{m}",{AccApp},"<>"),"")',
        f'=IFERROR(AVERAGEIFS({AccEnc},{AccM},"{m}",{AccEnc},"<>"),"")',
        f'=IFERROR(AVERAGEIFS({AccAbs},{AccM},"{m}",{AccAbs},"<>"),"")',
        f'=IF(ISNUMBER(E{r}),IF(E{r}<={thresh[m]},"PASS","FAIL"),"")'])
    sm.cell(row=r,column=5).number_format="0%"

# Repeatability block (plan rows 1-6 = repeatability, land in Trials rows 4-9 if entered in order)
rb=base+9
sm.cell(row=rb,column=1,value="REPEATABILITY  (block A — set the row range to your 6 repeatability trials)").font=F(bold=True,size=11,color=ACC)
hcell(rb+1,["","first row","last row"])
datarow(rb+2,["Trials rows",4,9])
sm.cell(rb+2,2).font=F(bold=True); sm.cell(rb+2,3).font=F(bold=True)
hcell(rb+3,["Metric","mean","SD (SEM)","MDC95 = 1.96·√2·SD"])
rep_metrics=[("app A0","A0"),("app f","f"),("app N","N"),("app R2n","R2n")]
for i,(cc,label) in enumerate(rep_metrics):
    r=rb+4+i
    rng=f'INDIRECT("Trials!{L(cc)}"&$B${rb+2}&":{L(cc)}"&$C${rb+2})'
    datarow(r,[label,
        f'=IFERROR(AVERAGE({rng}),"")',
        f'=IFERROR(STDEV({rng}),"")',
        f'=IFERROR(1.96*SQRT(2)*STDEV({rng}),"")'])
note=sm.cell(row=rb+9,column=1,value="Pass thresholds: f ≤5%, A0 ≤10%, N ≤15%, R2n ≤15%, area ≤20% mean absolute error vs encoder. MDC95 = smallest change one trial can resolve; compare A0 MDC95 to the 29° seen on participants.")
note.font=F(size=9,italic=True,color=GREY); sm.merge_cells(start_row=rb+9,start_column=1,end_row=rb+9,end_column=6); note.alignment=Alignment(wrap_text=True); sm.row_dimensions[rb+9].height=40
for j,w in enumerate([26,12,12,20,14,16],1): sm.column_dimensions[get_column_letter(j)].width=w

# ════════════════════════════════════════════════════ PROTOCOL
ps=wb.create_sheet("Protocol")
ps["A1"]="Protocol reference"; ps["A1"].font=F(bold=True,size=14)
ps["A2"]="The full plan the clean session is a first slice of. One setting per block; encoder trapezoid = 4 rings/file."; ps["A2"].font=F(size=9,italic=True,color=GREY)
PROTO=[["Block","Swept setting","Levels","Reps","Validates"],
 ["Z Test zero","Final Return Speed","200,150,100 (lift 50)","1","release floor that still rings"],
 ["A Repeatability","nothing","1","6-20","test-retest, all 7"],
 ["B Amplitude","Stretch Angle Range","37.5→5","1-3","A0, ωmax, ωmin, floor"],
 ["C Frequency","P term","0.43/1.0/1.7","1-3","f, ωmax vs filter"],
 ["D Decay","Derivative Term","0/small/smaller <0.01","1-3","R2n, N"],
 ["E Asymmetry","ZeroAngle offset","0/3/6°","1-3","P_asym (0° null)"]]
for i,rowv in enumerate(PROTO):
    for j,v in enumerate(rowv,1):
        cell=ps.cell(row=4+i,column=j,value=v); cell.border=box
        if i==0: cell.font=F(bold=True,color="FFFFFF"); cell.fill=fill(HEAD); cell.alignment=Alignment(horizontal="center",wrap_text=True)
        else: cell.font=F(size=10); cell.alignment=Alignment(vertical="center",wrap_text=True)
for j,w in enumerate([16,20,24,8,32],1): ps.column_dimensions[get_column_letter(j)].width=w

# ════════════════════════════════════════════════════ SESSION
ss=wb.create_sheet("Session")
ss["A1"]="Session log"; ss["A1"].font=F(bold=True,size=14)
for i,(k,v) in enumerate([("Session name",""),("Date",""),("Operator",""),("Phone / iOS",""),
      ("App build (COMMIT from app)",""),("Algorithm version (manifest)",""),("Rig",""),
      ("Swing plane","horizontal"),("Zero reference",""),("Encoder decoded folder",""),("Notes","")]):
    a=ss.cell(row=3+i,column=1,value=k); a.font=F(bold=True,size=10); a.border=box; a.fill=fill(ACC_SOFT)
    b=ss.cell(row=3+i,column=2,value=v); b.border=box; b.font=F(size=10)
ss.column_dimensions["A"].width=32; ss.column_dimensions["B"].width=48

# ════════════════════════════════════════════════════ HOW TO USE
rs=wb.create_sheet("How to use"); wb.move_sheet("How to use", -(len(wb.sheetnames)-1))
rs["A1"]="Bench Experiment Tracker"; rs["A1"].font=F(bold=True,size=16,color=INK)
for i,t in enumerate([
 "",
 "This workbook is the hub — the Bench Runner web tool is retired.",
 "",
 "AT THE BENCH",
 "  1. Fill the Session sheet once (date, operator, app COMMIT, encoder folder).",
 "  2. Follow the Clean Session Plan, top to bottom. Stretch mode only.",
 "  3. For each recording, add ONE row to Trials: identity, the rig settings you used,",
 "     crossings + settle seconds, and the app's readout. f and tau fill themselves.",
 "  4. The join that matters: Encoder file + segment. Write it the instant you record —",
 "     the clocks are independent, so it cannot be reconstructed later.",
 "",
 "AFTER THE BENCH",
 "  5. Run:  python bench_rescore.py --sheet <this file>   (fills the ENCODER TRUTH columns",
 "     by matching Phone JSONL file, scoring the encoder ring and re-scoring the phone",
 "     through the same function with the release supplied).",
 "  6. Read the Accuracy and Summary sheets — they compute themselves.",
 "",
 "WHAT THE SHEET ANALYSES",
 "  Trials  — f and tau from your counts.",
 "  Accuracy— app vs encoder error, per metric, per trial.",
 "  Summary — coverage (incl. app refusal rate), mean |%error| per metric with pass/fail,",
 "            and repeatability SEM + MDC95 for the block-A trials.",
 "",
 "PILOT FINDING (2026-09-11): sensor caught 35/35 swings; app refused 22/35; frequency",
 "accurate to ~4%; amplitude/velocity errors real but unmappable in that batch (manual",
 "record button, mixed manual/stretch, independent clocks). This clean session fixes that.",
]):
    cell=rs.cell(row=3+i,column=1,value=t)
    cell.font=F(bold=t in ("AT THE BENCH","AFTER THE BENCH","WHAT THE SHEET ANALYSES") or t.startswith("PILOT") or t.startswith("This workbook"),size=11,color=INK)
rs.column_dimensions["A"].width=104; rs.sheet_view.showGridLines=False

wb.save(OUT); print("wrote",OUT,"| sheets:",wb.sheetnames)
